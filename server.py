from __future__ import annotations

import asyncio
import base64
import json
import os
import re
import shutil
import sqlite3
import subprocess
import tempfile
import time
import uuid
from contextlib import asynccontextmanager, contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterator
from urllib.parse import quote

from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response
from pydantic import BaseModel
from pypdf import PdfReader
from websocket import create_connection

ROOT = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("RESUME_LAYOUT_DATA_DIR", str(ROOT / "data"))).resolve()
PDF_DIR = DATA_DIR / "pdfs"
TMP_DIR = DATA_DIR / "tmp"
DB_PATH = DATA_DIR / "resume.db"
PID_PATH = DATA_DIR / "server.pid"
APP_ID = "resume-layout-studio"
for directory in (DATA_DIR, PDF_DIR, TMP_DIR):
    directory.mkdir(parents=True, exist_ok=True)

@asynccontextmanager
async def lifespan(_: FastAPI):
    init_db()
    PID_PATH.write_text(str(os.getpid()), encoding="ascii")
    try:
        yield
    finally:
        if PID_PATH.exists() and PID_PATH.read_text(encoding="ascii").strip() == str(os.getpid()):
            PID_PATH.unlink()

app = FastAPI(title="简历排版工作台", version="1.0.0", lifespan=lifespan)

class ResumeBody(BaseModel):
    data: dict[str, Any]

class PdfBody(BaseModel):
    pages_html: str

def init_db() -> None:
    with db_connection() as connection:
        connection.execute("""
            CREATE TABLE IF NOT EXISTS resumes (
                id TEXT PRIMARY KEY, name TEXT NOT NULL, data_json TEXT NOT NULL,
                pdf_name TEXT, updated_at TEXT NOT NULL
            )
        """)


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def make_id(prefix: str) -> str:
    return f"{prefix}_{uuid.uuid4().hex[:12]}"


def safe_id(value: str) -> str:
    if not re.fullmatch(r"[A-Za-z0-9_-]{3,80}", value or ""):
        raise HTTPException(status_code=400, detail="无效的标识")
    return value


@contextmanager
def db_connection() -> Iterator[sqlite3.Connection]:
    connection = sqlite3.connect(DB_PATH, timeout=10)
    connection.row_factory = sqlite3.Row
    try:
        yield connection
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()

def normalize_resume(raw: dict[str, Any] | None, fallback_name: str = "未命名简历") -> dict[str, Any]:
    source = raw if isinstance(raw, dict) else {}
    basic_source = source.get("basic") if isinstance(source.get("basic"), dict) else {}
    basic = {
        "name": str(basic_source.get("name") or ""),
        "email": str(basic_source.get("email") or ""),
        "phone": str(basic_source.get("phone") or ""),
        "location": str(basic_source.get("location") or basic_source.get("city") or ""),
        "target": str(basic_source.get("target") or ""),
        "extra": str(basic_source.get("extra") or ""),
        "photo": str(basic_source.get("photo") or ""),
        "customFields": [],
    }
    custom_fields = basic_source.get("customFields")
    if not isinstance(custom_fields, list):
        custom_fields = basic_source.get("fields") or []
    if isinstance(custom_fields, list):
        basic["customFields"] = [
            {"label": str(field.get("label") or ""), "value": str(field.get("value") or "")}
            for field in custom_fields
            if isinstance(field, dict)
        ]

    sections: list[dict[str, Any]] = []
    for section_index, section in enumerate(source.get("sections") or []):
        if not isinstance(section, dict):
            continue
        items: list[dict[str, Any]] = []
        for item_index, item in enumerate(section.get("items") or []):
            if not isinstance(item, dict):
                continue
            bullets = item.get("bullets") if isinstance(item.get("bullets"), list) else []
            items.append(
                {
                    "id": str(item.get("id") or make_id("item")),
                    "title": str(item.get("title") or ""),
                    "subtitle": str(item.get("subtitle") or ""),
                    "date": str(item.get("date") or ""),
                    "summary": str(item.get("summary") or ""),
                    "bullets": [str(value) for value in bullets if str(value).strip()],
                }
            )
        sections.append(
            {
                "id": str(section.get("id") or f"section_{section_index + 1}"),
                "title": str(section.get("title") or "未命名模块"),
                "items": items,
            }
        )

    style_defaults = {
        "accent": "#2563eb",
        "font": "yahei",
        "titleFont": "simsun",
        "bodySize": 10.3,
        "lineHeight": 1.68,
        "marginX": 11,
        "marginY": 11,
        "sectionGap": 5,
        "entryGap": 4,
        "bulletIndent": 0,
        "photoSize": 24,
        "showPhoto": True,
    }
    if isinstance(source.get("style"), dict):
        style_defaults.update(source["style"])
    if str(style_defaults.get("accent") or "").lower() == "#ed5a0a":
        style_defaults["accent"] = "#2563eb"

    return {
        "id": str(source.get("id") or make_id("resume")),
        "name": str(source.get("name") or fallback_name),
        "templateId": str(source.get("templateId") or "classic_orange"),
        "basic": basic,
        "sections": sections,
        "style": style_defaults,
        "updatedAt": now_iso(),
    }


def resume_row(row: sqlite3.Row, include_data: bool = False) -> dict[str, Any]:
    parsed_data = json.loads(row["data_json"])
    result = {
        "id": row["id"],
        "name": row["name"],
        "templateId": parsed_data.get("templateId", "classic_orange"),
        "updatedAt": row["updated_at"],
        "hasPdf": bool(row["pdf_name"]),
        "pdfUrl": f"/api/resumes/{row['id']}/pdf?t={row['updated_at']}" if row["pdf_name"] else None,
    }
    if include_data:
        result["data"] = parsed_data
    return result


def upsert_resume(data: dict[str, Any]) -> dict[str, Any]:
    normalized = normalize_resume(data, str(data.get("name") or "未命名简历"))
    normalized["id"] = safe_id(normalized["id"])
    normalized["updatedAt"] = now_iso()
    with db_connection() as connection:
        existing = connection.execute("SELECT pdf_name FROM resumes WHERE id = ?", (normalized["id"],)).fetchone()
        pdf_name = existing["pdf_name"] if existing else None
        connection.execute(
            """
            INSERT INTO resumes (id, name, data_json, pdf_name, updated_at)
            VALUES (?, ?, ?, ?, ?)
            ON CONFLICT(id) DO UPDATE SET
                name = excluded.name,
                data_json = excluded.data_json,
                updated_at = excluded.updated_at
            """,
            (
                normalized["id"],
                normalized["name"],
                json.dumps(normalized, ensure_ascii=False),
                pdf_name,
                normalized["updatedAt"],
            ),
        )
    return normalized

def extract_text_from_upload(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".json":
        return content.decode("utf-8-sig")
    if suffix in {".txt", ".md", ".markdown"}:
        for encoding in ("utf-8-sig", "gb18030", "utf-16"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise HTTPException(status_code=400, detail="无法识别文本文件编码")
    if suffix == ".docx":
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False, dir=TMP_DIR) as stream:
            stream.write(content)
            temp_path = Path(stream.name)
        try:
            document = Document(temp_path)
            values = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        values.append(" | ".join(cells))
            return "\n".join(values)
        finally:
            temp_path.unlink(missing_ok=True)
    if suffix == ".pdf":
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False, dir=TMP_DIR) as stream:
            stream.write(content)
            temp_path = Path(stream.name)
        try:
            reader = PdfReader(str(temp_path))
            return "\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
        finally:
            temp_path.unlink(missing_ok=True)
    raise HTTPException(status_code=400, detail="仅支持 JSON、Word、PDF、Markdown 和 TXT")


SECTION_NAMES = {
    "教育背景": "education",
    "教育经历": "education",
    "实习经历": "experience",
    "工作经历": "experience",
    "项目经历": "projects",
    "产品项目": "projects",
    "科研经历": "research",
    "专业能力": "skills",
    "技能清单": "skills",
    "技能": "skills",
    "证书荣誉": "awards",
    "获奖经历": "awards",
    "自我评价": "summary",
    "个人优势": "strengths",
}


def heuristic_resume(text: str, filename: str) -> dict[str, Any]:
    lines = [re.sub(r"\s+", " ", line).strip(" \t•·#*-") for line in text.splitlines()]
    lines = [line for line in lines if line]
    email_match = re.search(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", text)
    phone_match = re.search(r"(?<!\d)1[3-9]\d{9}(?!\d)", text)
    name = next((line for line in lines[:8] if 2 <= len(line) <= 8 and not re.search(r"\d|@|简历", line)), "")

    buckets: list[tuple[str, str, list[str]]] = []
    current_title = "个人经历"
    current_id = "profile"
    current_lines: list[str] = []
    for line in lines:
        clean_title = re.sub(r"[：:｜|\s]", "", line)
        matched = next((title for title in SECTION_NAMES if re.sub(r"\s", "", title) == clean_title), None)
        if matched:
            if current_lines:
                buckets.append((current_id, current_title, current_lines))
            current_title = matched
            current_id = SECTION_NAMES[matched]
            current_lines = []
        else:
            current_lines.append(line)
    if current_lines:
        buckets.append((current_id, current_title, current_lines))

    sections = []
    for section_id, title, section_lines in buckets:
        filtered = [line for line in section_lines if line not in {name, email_match.group(0) if email_match else "", phone_match.group(0) if phone_match else ""}]
        if not filtered:
            continue
        sections.append(
            {
                "id": f"{section_id}_{len(sections) + 1}",
                "title": title,
                "items": [
                    {
                        "id": make_id("item"),
                        "title": "",
                        "subtitle": "",
                        "date": "",
                        "summary": "",
                        "bullets": filtered,
                    }
                ],
            }
        )

    return normalize_resume(
        {
            "name": f"{name or Path(filename).stem}的简历",
            "basic": {
                "name": name,
                "email": email_match.group(0) if email_match else "",
                "phone": phone_match.group(0) if phone_match else "",
            },
            "sections": sections,
        },
        Path(filename).stem,
    )

def find_browser() -> Path | None:
    candidates = [
        Path(os.environ.get("BROWSER_PATH", "")),
        Path(os.environ.get("LOCALAPPDATA", "")) / "Microsoft/Edge/Application/msedge.exe",
        Path(os.environ.get("LOCALAPPDATA", "")) / "Google/Chrome/Application/chrome.exe",
        Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
        Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
        Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
    ]
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    located = shutil.which("msedge") or shutil.which("chrome")
    return Path(located) if located else None


def sanitize_pages_html(value: str) -> str:
    cleaned = re.sub(r"<script\b[^>]*>.*?</script>", "", value, flags=re.I | re.S)
    cleaned = re.sub(r"<iframe\b[^>]*>.*?</iframe>", "", cleaned, flags=re.I | re.S)
    cleaned = re.sub(r"\s+on[a-z]+\s*=\s*(['\"]).*?\1", "", cleaned, flags=re.I | re.S)
    cleaned = re.sub(r"javascript\s*:", "", cleaned, flags=re.I)
    return cleaned


def create_pdf(resume_id: str, pages_html: str) -> Path:
    browser = find_browser()
    if not browser:
        raise RuntimeError("没有找到 Edge 或 Chrome")
    output_path = PDF_DIR / f"{resume_id}.pdf"
    css = (ROOT / "styles.css").read_text(encoding="utf-8")
    document = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><style>{css}</style>
<style>
@page {{ size: A4; margin: 0; }}
html, body {{ margin: 0 !important; padding: 0 !important; background: white !important; }}
.resume-pages {{ display: block !important; }}
.resume-page {{ margin: 0 !important; box-shadow: none !important; transform: none !important; break-after: page; page-break-after: always; }}
.resume-page:last-child {{ break-after: auto; page-break-after: auto; }}
</style></head><body><main class="resume-pages">{sanitize_pages_html(pages_html)}</main></body></html>"""
    with tempfile.TemporaryDirectory(dir=TMP_DIR, ignore_cleanup_errors=True) as profile:
        process = subprocess.Popen(
            [str(browser), "--headless", "--no-first-run", "--no-default-browser-check",
             "--disable-background-networking", "--enable-automation",
             "--remote-debugging-address=127.0.0.1", "--remote-debugging-port=0",
             f"--user-data-dir={profile}"],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        connection = None
        try:
            port_file = Path(profile) / "DevToolsActivePort"
            deadline = time.monotonic() + 30
            while not port_file.exists():
                if time.monotonic() >= deadline:
                    raise RuntimeError("浏览器启动超时，请更新 Edge 或 Chrome 后重试")
                time.sleep(0.1)
            port, endpoint = port_file.read_text(encoding="utf-8").splitlines()[:2]
            connection = create_connection(f"ws://127.0.0.1:{port}{endpoint}", timeout=30, suppress_origin=True, http_no_proxy=["127.0.0.1"])
            request_id = 0

            def call(method: str, params: dict | None = None, session: str | None = None) -> dict:
                nonlocal request_id
                request_id += 1
                message = {"id": request_id, "method": method, "params": params or {}}
                if session:
                    message["sessionId"] = session
                connection.send(json.dumps(message))
                while True:
                    result = json.loads(connection.recv())
                    if result.get("id") != request_id:
                        continue
                    if "error" in result:
                        raise RuntimeError("浏览器渲染失败，请更新浏览器后重试")
                    return result.get("result", {})

            target = call("Target.createTarget", {"url": "about:blank"})["targetId"]
            session = call("Target.attachToTarget", {"targetId": target, "flatten": True})["sessionId"]
            call("Page.enable", session=session)
            frame = call("Page.getFrameTree", session=session)["frameTree"]["frame"]["id"]
            call("Page.setDocumentContent", {"frameId": frame, "html": document}, session)
            call("Runtime.evaluate", {
                "expression": "Promise.all([document.fonts.ready, ...Array.from(document.images, image => image.complete ? Promise.resolve() : new Promise(resolve => { image.onload = resolve; image.onerror = resolve; }))])",
                "awaitPromise": True,
            }, session)
            pdf = call("Page.printToPDF", {
                "printBackground": True, "preferCSSPageSize": True,
                "displayHeaderFooter": False, "paperWidth": 210 / 25.4, "paperHeight": 297 / 25.4,
                "marginTop": 0, "marginBottom": 0, "marginLeft": 0, "marginRight": 0,
            }, session)
            output_path.write_bytes(base64.b64decode(pdf["data"]))
        finally:
            if connection:
                try:
                    connection.send(json.dumps({"id": request_id + 1, "method": "Browser.close"}))
                except Exception:
                    pass
                connection.close()
            if process.poll() is None:
                try:
                    process.wait(timeout=5)
                except subprocess.TimeoutExpired:
                    process.terminate()
    if not output_path.exists() or output_path.stat().st_size < 1000:
        raise RuntimeError("浏览器没有成功生成 PDF")
    return output_path

@app.get("/")
def index() -> FileResponse:
    return FileResponse(ROOT / "index.html", media_type="text/html; charset=utf-8", headers={"Cache-Control": "no-store"})


@app.get("/app.js")
def javascript() -> FileResponse:
    return FileResponse(ROOT / "app.js", media_type="application/javascript; charset=utf-8", headers={"Cache-Control": "no-store"})


@app.get("/styles.css")
def stylesheet() -> FileResponse:
    return FileResponse(ROOT / "styles.css", media_type="text/css; charset=utf-8", headers={"Cache-Control": "no-store"})

@app.get("/api/health")
def health() -> dict[str, Any]:
    return {
        "ok": True, "app": APP_ID, "instance": str(ROOT), "pid": os.getpid(),
        "version": app.version, "pdf": {"available": bool(find_browser())}
    }

@app.get("/api/example")
def example_resume() -> dict[str, Any]:
    data = json.loads((ROOT / "examples" / "演示简历.json").read_text(encoding="utf-8"))
    data["id"] = make_id("resume")
    return {"ok": True, "data": normalize_resume(data)}

@app.get("/api/resumes/{resume_id}/json")
def download_json(resume_id: str) -> Response:
    data = get_resume(resume_id)["data"]
    filename = quote(data["name"] + ".json", safe="")
    return Response(
        json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8"),
        media_type="application/json",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


@app.get("/api/resumes")
def list_resumes() -> list[dict[str, Any]]:
    init_db()
    with db_connection() as connection:
        rows = connection.execute("SELECT * FROM resumes ORDER BY updated_at DESC").fetchall()
    return [resume_row(row) for row in rows]


@app.get("/api/resumes/{resume_id}")
def get_resume(resume_id: str) -> dict[str, Any]:
    safe_id(resume_id)
    with db_connection() as connection:
        row = connection.execute("SELECT * FROM resumes WHERE id = ?", (resume_id,)).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="简历不存在")
    return resume_row(row, include_data=True)


@app.post("/api/resumes")
def save_resume(body: ResumeBody) -> dict[str, Any]:
    data = upsert_resume(body.data)
    return {"ok": True, "data": data}


@app.post("/api/resumes/{resume_id}/duplicate")
def duplicate_resume(resume_id: str) -> dict[str, Any]:
    safe_id(resume_id)
    with db_connection() as connection:
        row = connection.execute("SELECT data_json FROM resumes WHERE id = ?", (resume_id,)).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="简历不存在")
    data = json.loads(row["data_json"])
    data["id"] = make_id("resume")
    data["name"] = f"{data.get('name') or '未命名简历'} 副本"
    saved = upsert_resume(data)
    return {"ok": True, "data": saved}


@app.delete("/api/resumes/{resume_id}")
def delete_resume(resume_id: str) -> dict[str, Any]:
    safe_id(resume_id)
    with db_connection() as connection:
        cursor = connection.execute("DELETE FROM resumes WHERE id = ?", (resume_id,))
    if cursor.rowcount == 0:
        raise HTTPException(status_code=404, detail="简历不存在")
    (PDF_DIR / f"{resume_id}.pdf").unlink(missing_ok=True)
    return {"ok": True}


@app.get("/api/resumes/{resume_id}/pdf")
def download_pdf(resume_id: str) -> FileResponse:
    safe_id(resume_id)
    path = PDF_DIR / f"{resume_id}.pdf"
    if not path.exists():
        raise HTTPException(status_code=404, detail="尚未导出 PDF")
    with db_connection() as connection:
        row = connection.execute("SELECT name FROM resumes WHERE id = ?", (resume_id,)).fetchone()
    name = row["name"] if row else "简历"
    filename = re.sub(r"[\\/:*?\"<>|]", "_", name) + ".pdf"
    return FileResponse(path, media_type="application/pdf", filename=filename)


@app.post("/api/resumes/{resume_id}/pdf")
async def export_pdf(resume_id: str, body: PdfBody) -> dict[str, Any]:
    safe_id(resume_id)
    if not body.pages_html.strip() or len(body.pages_html) > 20_000_000:
        raise HTTPException(status_code=400, detail="预览内容无效或过大")
    with db_connection() as connection:
        row = connection.execute("SELECT id FROM resumes WHERE id = ?", (resume_id,)).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="请先保存简历")
    try:
        path = await asyncio.to_thread(create_pdf, resume_id, body.pages_html)
    except Exception as error:
        raise HTTPException(status_code=500, detail=f"PDF 生成失败：{error}") from error
    updated_at = now_iso()
    with db_connection() as connection:
        connection.execute(
            "UPDATE resumes SET pdf_name = ?, updated_at = ? WHERE id = ?",
            (path.name, updated_at, resume_id),
        )
    return {"ok": True, "url": f"/api/resumes/{resume_id}/pdf?t={updated_at}"}

@app.post("/api/import")
async def import_resume(file: UploadFile = File(...)) -> dict[str, Any]:
    content = await file.read(15 * 1024 * 1024 + 1)
    if len(content) > 15 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="文件不能超过 15 MB")
    filename = file.filename or "简历.txt"
    if Path(filename).suffix.lower() == ".json":
        try:
            parsed = json.loads(content.decode("utf-8-sig"))
        except (ValueError, UnicodeDecodeError) as error:
            raise HTTPException(status_code=400, detail="JSON 文件无效") from error
        if not isinstance(parsed, dict) or not isinstance(parsed.get("basic"), dict) or not isinstance(parsed.get("sections"), list):
            raise HTTPException(status_code=400, detail="JSON 需要包含 basic 对象和 sections 数组，请参考演示文件")
        data = normalize_resume(parsed, Path(filename).stem)
        data["id"] = make_id("resume")
        return {"ok": True, "data": data}
    try:
        text = extract_text_from_upload(filename, content)
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(status_code=400, detail="文件无法读取，请检查文件是否损坏或加密") from error
    if not text.strip():
        raise HTTPException(status_code=400, detail="没有提取到文字，扫描版 PDF 请先转为文字版")
    return {"ok": True, "data": heuristic_resume(text, filename), "notice": "已提取文字，请检查并整理条目、日期和模块"}

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=int(os.environ.get("RESUME_LAYOUT_PORT", "4180")), log_level="info")
